from hashlib import sha256
from pathlib import Path
from typing import Iterator

from docx import Document as DocxDocument
from pypdf import PdfReader

from app.ingestion.chunking import chunk_text
from app.schemas.documents import CitationLocation, Document, DocumentChunk, Modality


SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


def load_document(
    path: str | Path, *, max_chars: int = 1200, overlap: int = 120
) -> tuple[Document, list[DocumentChunk]]:
    """Load a supported file and return a normalized document plus chunks."""
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document extension: {suffix or '<none>'}")
    if not file_path.is_file():
        raise FileNotFoundError(file_path)

    document = Document(
        source=file_path.name,
        title=file_path.stem,
        modality=_modality_for(suffix),
        checksum=_checksum(file_path),
    )
    if overlap >= max_chars:
        overlap = 0
    chunks = list(_chunks(document, _extract(file_path, suffix), max_chars=max_chars, overlap=overlap))
    return document, chunks


def _extract(path: Path, suffix: str) -> Iterator[tuple[str, CitationLocation]]:
    if suffix == ".pdf":
        for page_number, page in enumerate(PdfReader(str(path)).pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                yield text, CitationLocation(page=page_number)
        return
    if suffix == ".docx":
        text = "\n".join(p.text.strip() for p in DocxDocument(str(path)).paragraphs if p.text.strip())
        if text:
            yield text, CitationLocation()
        return

    text = path.read_text(encoding="utf-8").strip()
    if text:
        yield text, CitationLocation()


def _chunks(
    document: Document,
    sections: Iterator[tuple[str, CitationLocation]],
    *,
    max_chars: int,
    overlap: int,
) -> Iterator[DocumentChunk]:
    ordinal = 0
    for section, citation in sections:
        for text in chunk_text(section, max_chars=max_chars, overlap=overlap):
            yield DocumentChunk(
                document_id=document.document_id,
                source=document.source,
                modality=document.modality,
                text=text,
                ordinal=ordinal,
                citation=citation,
            )
            ordinal += 1


def _modality_for(suffix: str) -> Modality:
    return {
        ".txt": Modality.TEXT,
        ".md": Modality.TEXT,
        ".markdown": Modality.TEXT,
        ".pdf": Modality.PDF,
        ".docx": Modality.DOCX,
    }[suffix]


def _checksum(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

